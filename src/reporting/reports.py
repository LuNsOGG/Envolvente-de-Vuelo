from pathlib import Path
from html import escape
from math import ceil
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from matplotlib import get_data_path
from matplotlib.figure import Figure
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from bootstrap import ROOT
from src.calculations import calculate,compare
from src.plotting.niu import FIGURES

FONTDIR=Path(get_data_path())/'fonts/ttf'

def text(value):return '' if value is None else str(value)
def number(value):return 'Pendiente' if value is None else (f'{value:.6g}' if isinstance(value,(float,int)) else str(value))
def chunks(seq,n):return [seq[i:i+n] for i in range(0,len(seq),n)]

def reference(p):
    src=p.get('fuente','')
    if '.jpg' in src:return '[F1] fotos'
    if 'govinfo' in src:return '[F3] FAR 2016'
    if not src:return 'Pendiente'
    return src

def equation_art(eq,directory,index):
    formulas={
       'Carga alar':r'$w_s=\frac{W}{S}$',
       'Pérdida EAS':r'$V_S=\frac{1}{k_v}\sqrt{\frac{2(W/S)}{\rho_0 C_{N+}}}$',
       'Intersección A':r'$q(A)\,C_+(A)/(W/S)=n_+$',
       'Relación de masas':r'$\mu_g=\frac{2(W/S)}{\rho\,\bar{c}\,a\,g}$',
       'Alivio de ráfaga':r'$K_g=\frac{0.88\mu_g}{5.3+\mu_g}$',
       'Ráfaga C':r'$\Delta n=\frac{K_g U_{de} V a}{498(W/S)},\quad n=1\pm\Delta n$',
       'Ráfaga D':r'$\Delta n=\frac{K_g U_{de} V a}{498(W/S)},\quad n=1\pm\Delta n$',
    }
    if eq['nombre'].endswith('_raw'):
        k={'Vc_A_raw':17,'Vd_A_raw':24,'Va_A_raw':15,'Vf_A_raw':11}[eq['nombre']]
        formula=rf'$V_{{\min,\mathrm{{base}}}}={k}\sqrt{{\max(n_1,3.8)\,(W/S)}}$'
    else:formula=formulas.get(eq['nombre'])
    if not formula:
        import textwrap
        formula='\n'.join(textwrap.wrap(eq.get('ecuación','Fórmula pendiente'),85))
    fig=Figure(figsize=(6.7,.55),dpi=220);fig.text(.02,.45,formula,fontsize=13 if not formula.startswith('$') else 16,va='center')
    path=directory/f'equation_{index:02}.png';fig.savefig(path,dpi=220,transparent=False)
    return path

def analysis_pages(project,case,result,plots,directory):
    if not result.get('requires_thesis',True):
        from .general_report import general_analysis_pages
        return general_analysis_pages(project,case,result,plots,directory)
    meta=project['metadata'];vals=result['values'];pages=[]
    pages.append(dict(title='Envolvente de vuelo Cessna Centurión',subtitle='Tarea 04 segundo parcial',cover=True,
        paragraphs=[f'Nombre: {meta.get("student") or "[Nombre editable]"}',f'Materia: {meta.get("subject") or "[Materia editable]"}',
                    f'Docente: {meta.get("teacher") or "[Docente editable]"}',f'Fecha: {meta.get("date","")}',
                    'Condición: '+case['name'],
                    'Análisis académico reproducible. La tesis no está disponible y la variante de Cessna todavía no se ha confirmado. El ejemplo con hipótesis permite revisar el método, pero no cierra la identificación de la aeronave ni demuestra cumplimiento normativo.']))
    pages.append(dict(title='Objetivo y alcance del análisis',paragraphs=[
        'Calcular y documentar carga alar, velocidades, cargas de maniobra y de ráfaga, sus intersecciones y la frontera combinada. Las tablas, curvas, sustituciones y resultados se generan a partir del mismo objeto de cálculo que utiliza la aplicación.',
        'Aeronave solicitada: Cessna Centurión. Modelo, variante, año, plazas, pesos, centro de gravedad, configuración y condición de la tesis: pendientes. El título de una aeronave cuatriplaza no permite inferir que los datos pertenezcan a cualquier Cessna 210; se necesita la tesis para resolverlo.',
        'Falta el archivo de la tesis Diseño de la caja de torsión de una aeronave cuatriplaza, especialmente las páginas de características de la aeronave, capítulo 1.9 Envolvente de vuelo y sus referencias normativas. Su ausencia bloquea la comparación numérica contra la tesis y la confirmación de la edición histórica utilizada.',
        'El libro adjunto es Airframe Structural Design, Michael Chun-Yung Niu. Su capítulo 3 comienza en visor 28, impresa 21. Se localizaron las figuras por número y pie de figura, conservando ambas paginaciones. La portada dice segunda edición; la página legal del escaneo indica copyright 1988 y segunda impresión enero de 1989. Se conserva esa discrepancia bibliográfica.',
        'Se consulta FAR 23 de 2016 como referencia oficial histórica independiente. Apéndice A y subparte C se implementan por separado. El primero usa n2=-0.5 n1 y factores n3/n4 de gráficos A1/A2; la subparte C normal usa el criterio de §23.337 y cargas de §23.341. No se atribuye esta edición a la tesis sin evidencia.',
        'Método de esta condición: '+case['method']+'. Los diagramas son estimaciones cuasiestáticas. No constituyen una certificación de aeronavegabilidad ni sustituyen datos de operación del manual de vuelo.']))
    for i,rows in enumerate(chunks(result['parameters'],7),1):
        data=[['Símbolo y descripción','Valor y unidad','Estado','Fuente y página','Observación']]
        for row in rows:
            data.append([row['símbolo']+'\n'+row['descripción'],number(row['valor'])+'\n'+row['unidad'],row['estado'],reference(row)+'\n'+row['página'],row['observaciones']])
        pages.append(dict(title=f'Trazabilidad de parámetros {i}',table=data,widths=[1.38,1.,1.,1.15,2.27],paragraphs=[
            'Documentado significa visible en la fuente indicada, no confirmado para una variante de Cessna. Las revisiones conservan lectura, unidad y referencia original en el JSON y en parametros.csv. La referencia F1 identifica las fotografías; F3 identifica FAR 23 de 2016.']))
    pages.append(dict(title='Hipótesis y datos pendientes',paragraphs=result['assumptions']+[('Cálculo bloqueado: '+x) for x in result['blocked']] or ['No hay hipótesis adicionales registradas.']))
    for i,eqs in enumerate(chunks(result['equations'],4),1):
        items=[]
        for eq in eqs:
            index=result['equations'].index(eq)
            items.append(dict(heading=eq['nombre'],image=equation_art(eq,directory,index),
                              body='Sustitución: '+eq['sustitución']+' = '+number(eq['resultado'])+' '+eq['unidad']+'.',source=eq['fuente']))
        pages.append(dict(title=f'Ecuaciones y sustituciones {i}',equations=items))
    pages.append(dict(title='Variables unidades y construcción de las fronteras',paragraphs=[
        'W es peso como fuerza en lbf; S, superficie en ft²; rho, densidad local en slug/ft³; rho0, densidad estándar de referencia EAS; c, cuerda geométrica media en ft; a, pendiente de fuerza normal por radián; g, gravedad en ft/s². La conversión de nudos a ft/s utiliza 1852/0.3048/3600. La carga alar es presión: lbf/ft².',
        'Una pendiente indicada por grado se multiplica por 180/pi para obtener el valor por radián. lbm/ft³ y kg/m³ se convierten a slug/ft³; la unidad ambigua lb/ft³ se rechaza. La aceleración g convierte peso a masa dentro de mu_g. No se trata lbf como lbm. Unidades CAS o TAS se rechazan en los campos EAS; TAS puede obtenerse por Ve/sqrt(rho/rho0) dentro de la aproximación incompresible.',
        'Se usa el coeficiente de fuerza normal CN para las curvas de pérdida; no se sustituye silenciosamente por CL. La pendiente del ala solo puede sustituir la pendiente normal de la aeronave bajo las condiciones de separación de cargas de cola que indica §23.341. Falta comprobar esas condiciones para el avión real.',
        'Maniobras de subparte C: rama superior min(CN+ q S/W,n+). Rama inferior max(CN- q S/W,n_límite(V)). En categoría normal el límite negativo sube linealmente desde n- en Vc hasta cero en Vd. Apéndice A conserva n2 en Vd, según A4, y usa los puntos A-C-D y G-F-E cuando C/F gobiernan.',
        'Para subparte C se calculan las ráfagas positivas y negativas alrededor de n=1. Entre Vc y Vd se interpolan linealmente los factores completos, conforme a §23.333(c)(2)(ii). Interpolar la intensidad de ráfaga y después multiplicarla por V produciría otra curva y no se hace.',
        'El combinado de subparte C conserva los extremos de las condiciones de maniobra y ráfaga a cada velocidad y se limita mediante las curvas estáticas de CN. Las rectas de ráfaga completas se mantienen visibles para distinguir la demanda lineal de esa limitación. Es una construcción académica cuasiestática; una ráfaga dinámica requeriría otro análisis. No se recorta una carga de ráfaga por superar el factor de maniobra.',
        'El modo de apéndice A sigue O-A-C-D-E-F-G-O de A4 con n3 y n4 introducidos desde A1/A2. No se digitalizan automáticamente esos gráficos ni se aplica una igualdad n3=n1 o n4=n2 sin lectura documentada. La definición exacta que adopta la tesis permanece pendiente.']))
    points=[['Punto','V kt EAS','n']]+[[p['punto'],number(p['V']),number(p['n'])] for p in result['points']]
    ints=[['Intersección','V kt EAS','n']]+[[p['tipo'],number(p['V']),number(p['n'])] for p in result['intersections']]
    pages.append(dict(title='Puntos característicos e intersecciones',table=points,widths=[3.4,1.7,1.7],second_table=ints,
                     paragraphs=['Las raíces se obtienen de ecuaciones cuadráticas por tramo y se incluyen en la malla exportada. Va seleccionada y la intersección A se identifican por separado. Todas las filas completas de curvas están en curvas.csv.']))
    for kind,title in [('man','Diagrama de maniobras'),('gust','Diagrama de ráfagas'),('combined','Envolvente combinada')]:
        pages.append(dict(title=title,image=plots[kind+'_png'],paragraphs=[
            'Condición: '+case['name']+'. Las velocidades están en kt EAS y los factores son adimensionales. Consulte hipótesis y verificaciones antes de interpretar el contorno.',
            'La imagen procede de las mismas listas numéricas de curvas.csv. Los coeficientes, límites y velocidades no se alteran para modificar la apariencia de la gráfica.']))
    for i,cks in enumerate(chunks(result['checks'],8),1):
        pages.append(dict(title=f'Revisión de congruencia {i}',paragraphs=[f'{x["estado"]} | {x["código"]}. {x["detalle"]}' for x in cks]))
    rows=[['Variable','Original fotos','Ejemplo revisado','Peso menor']]
    results={c['id']:calculate(c) for c in project['cases']}
    ov=results.get('original_fotos',result)['values']
    nw_note=(f'En el caso original de fotos se leen Wn={number(ov.get("Wn_photo"))}, Wu={number(ov.get("Wu_photo"))}, '
             f'nn={number(ov.get("nn_photo"))} y nu={number(ov.get("nu_photo"))}. Sus productos nW son '
             f'{number(ov.get("nW_normal_photo"))} y {number(ov.get("nW_utility_photo"))} lbf respectivamente. '
             'A23.5 exige analizar esa relación para varias categorías. El significado de los subíndices y las condiciones de cada peso requieren la tesis; no se confirma equivalencia multicategoría.')
    for key in ['W_S','Vc_A_raw','Vd_A_raw','Va_A_raw','Vs','V_A_stall','mu','Kg','n_gust_c_pos','n_gust_c_neg']:
        rows.append([key]+[number(results.get(cid,{}).get('values',{}).get(key)) for cid in ['original_fotos','revisado_ilustrativo','peso_menor']])
    pages.append(dict(title='Comparación entre condiciones',table=rows,widths=[2.1,1.5,1.6,1.6],paragraphs=[
        'El original solo confirma las operaciones posibles con sus datos. Las celdas pendientes no representan cero. La comparación de peso mantiene las velocidades seleccionadas y las demás hipótesis; no se interpreta automáticamente como otra categoría de certificación.',
        nw_note]))
    for i,changes in enumerate(chunks(case.get('changes',[]),8),1):
        data=[['Parámetro','Original','Revisado','Justificación']]
        for c in changes:data.append([c['parameter'],number(c.get('before'))+' '+c.get('before_unit',''),number(c.get('after'))+' '+c.get('after_unit',''),c.get('reason','')])
        pages.append(dict(title=f'Cambios y efecto sobre el cálculo {i}',table=data,widths=[.8,1.15,1.15,3.7],paragraphs=[
            'La tabla anterior compara los resultados obtenidos. Cambiar la unidad de a modifica directamente la pendiente convertida y las cargas de ráfaga; aclarar rho y c permite calcular mu y Kg; introducir CN negativo permite cerrar la rama inferior. Las selecciones Vc/Vd definen el intervalo y los puntos C/D. La modificación de criterio cambia la ley negativa y las condiciones normativas; no es una corrección confirmada de la tesis.']))
    discussion='Los cálculos completos están bloqueados por los datos indicados.'
    if 'n_gust_c_pos' in vals:
        discussion=(f'En esta condición W/S={vals["W_S"]:.6f} lbf/ft², Vs={vals["Vs"]:.3f} kt EAS y A={vals["V_A_stall"]:.3f} kt EAS. '
                    f'La relación de masas es {vals["mu"]:.5f} y Kg={vals["Kg"]:.6f}. En Vc={vals["Vc"]:g} kt EAS las ráfagas producen '
                    f'n+={vals["n_gust_c_pos"]:.5f} y n-={vals["n_gust_c_neg"]:.5f}. Estos son resultados del caso con sus hipótesis, no valores verificados de la aeronave.')
    pages.append(dict(title='Discusión y conclusiones',paragraphs=[discussion,
        'Las demandas de ráfaga pueden superar el límite seleccionado para maniobras. La estructura se debe investigar para las condiciones que gobiernen según el criterio aplicable; superar el límite de maniobra no es un error aritmético que deba borrarse. La aplicación conserva ambos conjuntos de cargas y muestra su relación.',
        'La reducción de peso aumenta la respuesta normalizada de ráfaga en la comparación ilustrativa, aun cuando cambia el alivio Kg. Los mínimos de velocidad del apéndice A se calculan con el peso y n1 seleccionados; no se confunden con velocidades de operación ni con la velocidad real de pérdida de una variante.',
        ('Si la cuerda introducida es la geométrica S/b, el alargamiento inferido S/c² es '+number(vals.get('AR_conditional'))+'. Consulte la verificación AR CONDICIONAL A23.1: un valor superior a 7 haría incompatible el apéndice A bajo esa interpretación. Si la cuerda fuera aerodinámica media, esa inferencia no sería válida. Se necesita la geometría real para resolverlo.'),
        'Para concluir la actividad con el procedimiento original se requiere la tesis, confirmar la variante y plazas, pesos y CG, la edición normativa, la geometría de A23.1, las unidades de rho/c/a, la pendiente normal o tratamiento separado de cola, el coeficiente normal negativo y las velocidades seleccionadas. Si se usa apéndice A también se necesita la lectura de n3/n4 de A1/A2.',
        'Los resultados permiten revisar ecuaciones, sustituciones, consistencia y gráficos. No se ha demostrado la resistencia de una estructura, el comportamiento dinámico ni el cumplimiento normativo de un avión. Los resultados solo deben utilizarse en el alcance académico declarado.']))
    pages.append(dict(title='Referencias y reproducibilidad',paragraphs=[
        '[F1] Fotografías adjuntas: 308aa1a8-3551-4e50-a2fb-7b7e27fd85c9.jpg y cd1a2819-dc32-4a24-b668-615ca1cc63e5.jpg. Lectura visual de símbolos y valores; celdas borrosas de velocidades seleccionadas no confirmadas.',
        '[F2] Niu, Michael Chun-Yung. Airframe Structural Design. Conmilit Press. PDF adjunto: portada rotulada segunda edición; página legal indica 1988 y segunda impresión 1989. Capítulo 3 Aircraft Loads, visor 28 a 96. Reporte independiente con localizadores por figura.',
        '[F3] Office of the Federal Register / FAA. 14 CFR Part 23, edición 1 de enero de 2016. Apéndice A, A23.1, A23.3, A23.5, A23.7 y A23.9; Tabla 1, p.342; figuras A1-A4, pp.343-345. Copia oficial guardada en sources/.',
        '[F4] 14 CFR §§23.333, 23.335, 23.337 y 23.341, edición 2016, pp.216-219. Ráfagas 50/25 ft/s para 0 a 20000 ft y reducción lineal por altitud hasta 50000 ft. Pendiente por radián y velocidad equivalente en nudos.',
        'Enlace oficial del apéndice: https://www.govinfo.gov/content/pkg/CFR-2016-title14-vol1/pdf/CFR-2016-title14-vol1-part23-appA.pdf',
        'Enlace oficial de ráfagas: https://www.govinfo.gov/content/pkg/CFR-2016-title14-vol1/pdf/CFR-2016-title14-vol1-sec23-341.pdf',
        '[PENDIENTE] Tesis Diseño de la caja de torsión de una aeronave cuatriplaza. No disponible en la carpeta; no se atribuyen datos ni ecuaciones a páginas no consultadas.',
        'Reproducir con Ejecutar_calculos.cmd o python main.py --all. El proyecto.json exportado conserva la configuración utilizada; las huellas de fuentes se registran sin modificar los originales. Los CSV usan UTF-8 con BOM. Cada exportación genera una carpeta nueva.']))
    return pages

def niu_pages(project,images):
    meta=project['metadata']
    pages=[dict(title='Análisis de figuras de cargas de aeronaves',subtitle='Airframe Structural Design capítulo 3',cover=True,author=meta.get('student',''),paragraphs=[
        f'Nombre: {meta.get("student") or "[Nombre editable]"}',f'Materia: {meta.get("subject") or "[Materia editable]"}',
        f'Docente: {meta.get("teacher") or "[Docente editable]"}',f'Fecha: {meta.get("date","")}',
        'Se analizan y redibujan las figuras 3.1.3 a y b, 3.3.1 a 3.3.4, 3.4.8, 3.4.9, 3.4.11, 3.4.16 y 3.5.6. El objetivo es relacionar fuerzas, momentos y condiciones críticas con el diseño estructural.',
        'Fuente: Michael Chun-Yung Niu, Airframe Structural Design, Conmilit Press. El archivo recibido muestra segunda edición en la portada y datos legales de 1988 / segunda impresión de 1989. Se registra esa inconsistencia sin inferir otra edición.',
        'El capítulo Aircraft Loads comienza en visor 28, impresa 21. Los dibujos de este reporte son interpretaciones técnicas originales; sus simplificaciones se declaran en cada figura y las coordenadas no son datos de una aeronave real.'])]
    for f in FIGURES:
        pages.append(dict(title='Figura '+f['number']+' '+f['spanish'],image=images[(f['number'],'png')],image_height=3.5,
            paragraphs=[f'Título original: {f["title"]} Página impresa {f["page"]}; página del visor {f["viewer"]}.',
                        f['explanation'],f['symbols'],'Simplificación del redibujo: '+f['simplification']]))
    return pages

def render_pdf(pages,path):
    for name,filename in [('App','DejaVuSans.ttf'),('AppBold','DejaVuSans-Bold.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(FONTDIR/filename)))
    styles=getSampleStyleSheet()
    body=ParagraphStyle('BodyApp',fontName='App',fontSize=10,leading=14,spaceAfter=9)
    small=ParagraphStyle('SmallApp',parent=body,fontSize=8,leading=10,spaceAfter=0)
    h=ParagraphStyle('HeadingApp',parent=body,fontName='AppBold',fontSize=17,leading=22,spaceAfter=18,textColor=colors.HexColor('#111111'))
    title=ParagraphStyle('TitleApp',parent=h,fontSize=27,leading=33,spaceAfter=25)
    def para(s,style=body):return Paragraph(escape(text(s)).replace('\n','<br/>'),style)
    def table(data,widths=None):
        total=6.8*72
        widths=[x*72 for x in widths] if widths else [total/len(data[0])]*len(data[0])
        tab=Table([[para(x,small) for x in row] for row in data],colWidths=widths,repeatRows=1,hAlign='LEFT')
        tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e5e5e5')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f5f5f5')]),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#cccccc')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        return tab
    story=[]
    for i,p in enumerate(pages):
        if i:story.append(PageBreak())
        if p.get('cover'):story.append(Spacer(1,50))
        story.append(para(p['title'],title if p.get('cover') else h))
        if p.get('subtitle'):story.append(para(p['subtitle'],h))
        if p.get('image'):
            im=Image(str(p['image']));im._restrictSize(6.8*72,p.get('image_height',4.35)*72);story.extend([im,Spacer(1,12)])
        for s in p.get('paragraphs',[]):story.append(para(s))
        if p.get('table'):story.extend([table(p['table'],p.get('widths')),Spacer(1,12)])
        if p.get('second_table') and len(p['second_table'])>1:story.extend([para('Intersecciones calculadas',body),table(p['second_table'],p.get('widths'))])
        for eq in p.get('equations',[]):
            elements=[para(eq['heading'],ParagraphStyle('EqHeading',parent=body,fontName='AppBold'))]
            if eq.get('image'):elements.append(Image(str(eq['image']),width=6.7*72,height=.55*72))
            elements.extend([para(eq['body']),para(eq['source'],ParagraphStyle('EqSource',parent=body,fontSize=8,leading=11)),Spacer(1,8)])
            story.append(KeepTogether(elements))
    def footer(canvas,doc):
        canvas.setFont('App',8);canvas.setFillColor(colors.HexColor('#666666'));canvas.drawString(54,27,pages[0].get('footer','Tarea 04 · Análisis académico'));canvas.drawRightString(letter[0]-54,27,str(doc.page))
    doc=SimpleDocTemplate(str(path),pagesize=letter,leftMargin=61.2,rightMargin=61.2,topMargin=48,bottomMargin=47,title=pages[0]['title'],author=pages[0].get('author',''))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

def render_docx(pages,path):
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.line_spacing=1.08
    for s in ['Title','Heading 1','Heading 2']:
        doc.styles[s].font.name='Calibri';doc.styles[s].font.color.rgb=RGBColor.from_string('000000')
    doc.core_properties.author=pages[0].get('author','');doc.core_properties.title=pages[0]['title']
    def add_table(data,widths=None):
        tab=doc.add_table(rows=0,cols=len(data[0]));tab.autofit=False
        if widths:
            for col,w in zip(tab.columns,widths):col.width=Inches(w)
        for idx,row in enumerate(data):
            cells=tab.add_row().cells
            for j,(cell,value) in enumerate(zip(cells,row)):
                if widths:cell.width=Inches(widths[j])
                cell.text=text(value)
                for pp in cell.paragraphs:
                    pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4)
                    for run in pp.runs:run.font.size=Pt(8.5)
                props=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E5E5E5' if idx==0 else ('F5F5F5' if idx%2 else 'FFFFFF'));props.append(shd)
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                props.append(borders)
            if idx==0:
                tr=tab.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');tr.append(repeat)
        return tab
    for i,p in enumerate(pages):
        if i:doc.add_page_break()
        doc.add_paragraph(p['title'],'Title' if p.get('cover') else 'Heading 1')
        if p.get('subtitle'):doc.add_paragraph(p['subtitle'],'Subtitle')
        if p.get('image'):
            from PIL import Image as PILImage
            im=PILImage.open(p['image']);width=min(6.65,p.get('image_height',4.35)*im.width/im.height)
            doc.add_picture(str(p['image']),width=Inches(width))
        for s in p.get('paragraphs',[]):doc.add_paragraph(s)
        if p.get('table'):add_table(p['table'],p.get('widths'))
        if p.get('second_table') and len(p['second_table'])>1:
            doc.add_paragraph('Intersecciones calculadas','Heading 2');add_table(p['second_table'],p.get('widths'))
        for eq in p.get('equations',[]):
            doc.add_paragraph(eq['heading'],'Heading 2')
            if eq.get('image'):doc.add_picture(str(eq['image']),width=Inches(6.5))
            doc.add_paragraph(eq['body']);doc.add_paragraph(eq['source'])
    # Las fórmulas se insertan como imágenes; texto y tablas permanecen editables.
    doc.save(path)
